"""Parse-subprocess entry point (M4_PLAN.md 4.2, 14.2), sandboxed by RLIMIT_AS / RLIMIT_CPU on Linux.

Invoked by :mod:`semigraph.uploads.parse` as ``python -m semigraph.uploads.parse_worker <kind> <max_pages>`` with
the raw document bytes on stdin; it writes ONE JSON object to stdout and exits 0 on success, or writes
``{"error": code}`` and exits non-zero. The rlimits are set as this module's FIRST statements — before importing
any parser (pypdfium2 / pdfplumber / python-docx) and before reading stdin — and NEVER via ``preexec_fn`` (unsafe
in a multi-threaded parent: the child can deadlock before exec). ``-m`` itself imports ``semigraph`` and
``semigraph.uploads`` first; both are import-light (stdlib only), so nothing heavy runs before the limits are set.
"""

import sys

_RLIMIT_AS_BYTES = 1 * 1024 * 1024 * 1024        # 1 GiB
_RLIMIT_CPU_SECONDS = 120


def _lower_rlimit(kind: int, soft: int, resource_mod) -> None:
    _, hard = resource_mod.getrlimit(kind)
    cap = soft if hard == resource_mod.RLIM_INFINITY else min(soft, hard)
    resource_mod.setrlimit(kind, (cap, hard))


if sys.platform.startswith("linux"):                # pragma: no cover (exercised by the Linux serve-shipped CI job)
    import resource as _resource

    _lower_rlimit(_resource.RLIMIT_AS, _RLIMIT_AS_BYTES, _resource)
    _lower_rlimit(_resource.RLIMIT_CPU, _RLIMIT_CPU_SECONDS, _resource)

import json          # noqa: E402  (imported only after the rlimit guard above, on purpose)
import logging       # noqa: E402
import re            # noqa: E402

logger = logging.getLogger("semigraph.uploads.parse_worker")

PARSE_ERROR_CODES = ("parse_failed", "timeout", "too_large", "scanned", "empty", "too_many_pages")
_MIN_TEXT_CHARS_PER_PAGE = 200
DEFAULT_BODY_PT = 12.0
_DOCX_CHARS_PER_ESTIMATED_PAGE = 3000
_HEADING_TAGS = ("h1", "h2", "h3", "h4", "h5", "h6")
_TABLE_CELL_TAGS = ("td", "th")
_BLOCK_TAGS = _HEADING_TAGS + _TABLE_CELL_TAGS + ("p", "li", "div")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


class _ParseWorkerError(Exception):
    def __init__(self, code: str) -> None:
        if code not in PARSE_ERROR_CODES:
            raise ValueError(f"unknown parse error code {code!r}")
        super().__init__(code)
        self.code = code


def _merge_lines_into_blocks(lines: list[tuple[str, float, bool]], page: int) -> list[dict]:
    """Consecutive lines of the SAME style (rounded size, same bold) become one paragraph block, joined by a
    single space (never a newline: a newline would read as a sentence break downstream, section 14.2/4.2)."""
    merged: list[list] = []
    for text, size, bold in lines:
        if merged and merged[-1][2] == bold and abs(merged[-1][1] - size) < 0.5:
            merged[-1][0] = f"{merged[-1][0]} {text}"
        else:
            merged.append([text, size, bold])
    return [{"text": t, "page": page, "size": s, "bold": b, "kind_hint": "paragraph"} for t, s, b in merged]


# --------------------------------------------------------------------------
# PDF (pypdfium2 raw API, pdfplumber cross-check)
# --------------------------------------------------------------------------

def _pdf_page_lines(text_page_raw, praw, ctypes) -> list[tuple[str, float, bool]]:
    """One (text, effective_size, bold) line per baseline-y run of characters on the page.

    Effective size is ``FPDFText_GetMatrix`` scale (``matrix.d``) times ``FPDFText_GetFontSize`` (a Word-exported
    PDF can report a constant nominal size of 1.0 with the real size folded into the matrix; multiplying the two
    is correct either way — verified). Bold is ``FPDFText_GetFontWeight() >= 600`` OR a "Bold" font name (the
    standard-14 fonts carry no weight in their descriptor and always report 0 — verified). Characters are read in
    native index order (a geometric re-sort scrambles text — verified bug, section 3)."""
    n = praw.FPDFText_CountChars(text_page_raw)
    if n <= 0:
        return []
    x, y = ctypes.c_double(), ctypes.c_double()
    matrix = praw.FS_MATRIX()
    name_buf = ctypes.create_string_buffer(256)
    flags = ctypes.c_int(0)
    lines: list[tuple[str, float, bool]] = []
    cur_chars: list[str] = []
    cur_y = cur_size = None
    cur_bold = False
    for i in range(n):
        code = praw.FPDFText_GetUnicode(text_page_raw, i)
        praw.FPDFText_GetCharOrigin(text_page_raw, i, ctypes.byref(x), ctypes.byref(y))
        praw.FPDFText_GetMatrix(text_page_raw, i, matrix)
        font_size = praw.FPDFText_GetFontSize(text_page_raw, i) or 0.0
        eff_size = (matrix.d or 1.0) * font_size or 12.0
        weight = praw.FPDFText_GetFontWeight(text_page_raw, i)
        name_len = praw.FPDFText_GetFontInfo(text_page_raw, i, name_buf, 256, ctypes.byref(flags))
        font_name = name_buf.raw[:name_len].decode("utf-8", "replace").rstrip("\x00").lower()
        bold = weight >= 600 or "bold" in font_name
        is_new_line = cur_y is None or abs(y.value - cur_y) > max(eff_size, 1.0) * 0.4
        if is_new_line:
            _flush_pdf_line(cur_chars, cur_size, cur_bold, lines)
            cur_chars, cur_y, cur_size, cur_bold = [], y.value, eff_size, bold
        try:
            cur_chars.append(chr(code) if code else "")
        except (ValueError, OverflowError):
            cur_chars.append("")
    _flush_pdf_line(cur_chars, cur_size, cur_bold, lines)
    return lines


def _flush_pdf_line(chars: list[str], size: float | None, bold: bool, out: list[tuple[str, float, bool]]) -> None:
    text = "".join(chars).strip()
    if text:
        out.append((text, size or DEFAULT_BODY_PT, bold))


def _pdfium_extract(data: bytes, n_pages: int) -> tuple[list[dict], float]:
    import ctypes as _ctypes

    import pypdfium2 as pdfium
    import pypdfium2.raw as praw

    pdf = pdfium.PdfDocument(data)
    try:
        blocks: list[dict] = []
        total_chars = 0
        for page_no in range(n_pages):
            page = pdf[page_no]
            textpage = page.get_textpage()
            try:
                lines = _pdf_page_lines(textpage.raw, praw, _ctypes)
            finally:
                textpage.close()
                page.close()
            total_chars += sum(len(t) for t, _, _ in lines)
            blocks.extend(_merge_lines_into_blocks(lines, page_no + 1))
        return blocks, (total_chars / n_pages if n_pages else 0.0)
    finally:
        pdf.close()


def _group_plumber_words(words: list[dict]) -> list[tuple[str, float, bool]]:
    rows: dict[int, list[dict]] = {}
    for w in words:
        rows.setdefault(round(w["top"]), []).append(w)
    lines = []
    for top in sorted(rows):
        ws = sorted(rows[top], key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in ws)
        size = float(ws[0].get("size") or DEFAULT_BODY_PT)
        bold = "bold" in str(ws[0].get("fontname") or "").lower()
        lines.append((text, size, bold))
    return lines


def _pdfplumber_extract(data: bytes, n_pages: int) -> tuple[list[dict], float]:
    import io

    import pdfplumber

    blocks: list[dict] = []
    total_chars = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(extra_attrs=["size", "fontname"])
            lines = _group_plumber_words(words)
            total_chars += sum(len(t) for t, _, _ in lines)
            blocks.extend(_merge_lines_into_blocks(lines, page_no))
    return blocks, (total_chars / n_pages if n_pages else 0.0)


def _parse_pdf(data: bytes, max_pages: int) -> dict:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
        n_pages = len(pdf)
        pdf.close()
    except Exception as exc:
        raise _ParseWorkerError("parse_failed") from exc
    if n_pages == 0:
        raise _ParseWorkerError("empty")
    if n_pages > max_pages:
        raise _ParseWorkerError("too_many_pages")

    blocks, chars_per_page = _pdfium_extract(data, n_pages)
    method, warnings = "pypdfium2", []
    if chars_per_page < _MIN_TEXT_CHARS_PER_PAGE:
        fallback_blocks, fallback_cpp = _pdfplumber_extract(data, n_pages)
        if fallback_cpp >= _MIN_TEXT_CHARS_PER_PAGE:
            blocks, chars_per_page, method = fallback_blocks, fallback_cpp, "pdfplumber"
            warnings.append("pypdfium2 yielded < 200 chars/page; used the pdfplumber cross-check")
        elif not blocks and not fallback_blocks:
            raise _ParseWorkerError("empty")
        else:
            raise _ParseWorkerError("scanned")
    return {"method": method, "pages": n_pages, "blocks": blocks, "chars_per_page": chars_per_page,
            "warnings": warnings}


# --------------------------------------------------------------------------
# DOCX (python-docx: iter_inner_content covers paragraphs AND tables)
# --------------------------------------------------------------------------

def _docx_paragraph_style(paragraph) -> tuple[bool, float, bool]:
    """(is_structural_heading, size, bold) of one paragraph: a real ``Heading N`` style, else the run-level
    bold/size fallback (a bold 12pt run still reports style "Normal" — verified python-docx quirk)."""
    style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
    if style_name.lower().replace(" ", "").startswith(("heading", "title")):
        return True, 18.0, True
    non_empty_runs = [run for run in paragraph.runs if run.text.strip()]
    bold = bool(non_empty_runs) and all(run.bold for run in non_empty_runs)
    size = next((run.font.size.pt for run in non_empty_runs if run.font.size), DEFAULT_BODY_PT)
    return False, float(size), bold


def _parse_docx(data: bytes) -> dict:
    import io

    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise _ParseWorkerError("parse_failed") from exc

    blocks: list[dict] = []
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            is_heading, size, bold = _docx_paragraph_style(item)
            kind_hint = "heading_style" if is_heading else "paragraph"
            blocks.append({"text": text, "page": 1, "size": size, "bold": bold, "kind_hint": kind_hint})
        elif isinstance(item, Table):
            blocks.extend(_docx_table_blocks(item))
    if not blocks:
        raise _ParseWorkerError("empty")
    total_chars = sum(len(b["text"]) for b in blocks)
    return {"method": "python-docx", "pages": 1, "blocks": blocks, "chars_per_page": float(total_chars),
            "warnings": [], "_estimated_pages": max(1, -(-total_chars // _DOCX_CHARS_PER_ESTIMATED_PAGE))}


def _docx_table_blocks(table) -> list[dict]:
    out = []
    for row in table.rows:
        for cell in row.cells:
            text = cell.text.strip()
            if text:
                out.append({"text": text, "page": 1, "size": DEFAULT_BODY_PT, "bold": False, "kind_hint": "table"})
    return out


# --------------------------------------------------------------------------
# HTML (stdlib html.parser)
# --------------------------------------------------------------------------

def _html_heading_size(tag: str) -> float:
    return 24.0 - (int(tag[1]) - 1) * 2.0


class _BlockHTMLParser:
    """A tiny block-level text extractor built on ``html.parser.HTMLParser``; script/style content is dropped."""

    def __init__(self) -> None:
        from html.parser import HTMLParser

        outer = self
        blocks: list[dict] = []
        state = {"skip": 0, "tag": None, "buf": []}

        class _Parser(HTMLParser):
            def handle_starttag(self, tag, attrs):
                if tag in ("script", "style", "title"):
                    state["skip"] += 1
                elif tag in _BLOCK_TAGS:
                    outer._flush(state, blocks)
                    state["tag"] = tag

            def handle_endtag(self, tag):
                if tag in ("script", "style", "title"):
                    state["skip"] = max(0, state["skip"] - 1)
                elif tag in _BLOCK_TAGS:
                    outer._flush(state, blocks)
                    state["tag"] = None

            def handle_data(self, data):
                if state["skip"] == 0:
                    state["buf"].append(data)

        self._parser, self._blocks, self._state = _Parser(), blocks, state

    def _flush(self, state: dict, blocks: list[dict]) -> None:
        text = " ".join("".join(state["buf"]).split())
        state["buf"] = []
        if not text:
            return
        tag = state["tag"] or "p"
        if tag in _HEADING_TAGS:
            blocks.append({"text": text, "page": 1, "size": _html_heading_size(tag), "bold": True,
                          "kind_hint": "heading_html"})
        elif tag in _TABLE_CELL_TAGS:
            blocks.append({"text": text, "page": 1, "size": DEFAULT_BODY_PT, "bold": tag == "th",
                          "kind_hint": "table"})
        else:
            blocks.append({"text": text, "page": 1, "size": DEFAULT_BODY_PT, "bold": False, "kind_hint": "paragraph"})

    def parse(self, html_text: str) -> list[dict]:
        self._parser.feed(html_text)
        self._parser.close()
        self._flush(self._state, self._blocks)
        return self._blocks


def _parse_html(data: bytes) -> dict:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise _ParseWorkerError("parse_failed") from exc
    blocks = _BlockHTMLParser().parse(text)
    if not blocks:
        raise _ParseWorkerError("empty")
    total_chars = sum(len(b["text"]) for b in blocks)
    return {"method": "html.parser", "pages": 1, "blocks": blocks, "chars_per_page": float(total_chars),
            "warnings": []}


# --------------------------------------------------------------------------
# Markdown / plain text
# --------------------------------------------------------------------------

def _parse_md(data: bytes) -> dict:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise _ParseWorkerError("parse_failed") from exc
    blocks = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = _MD_HEADING_RE.match(line)
        if match:
            level = len(match.group(1))
            blocks.append({"text": match.group(2).strip(), "page": 1, "size": 24.0 - (level - 1) * 2.0,
                          "bold": True, "kind_hint": "heading_md"})
        else:
            blocks.append({"text": line, "page": 1, "size": DEFAULT_BODY_PT, "bold": False, "kind_hint": "paragraph"})
    if not blocks:
        raise _ParseWorkerError("empty")
    total_chars = sum(len(b["text"]) for b in blocks)
    return {"method": "text", "pages": 1, "blocks": blocks, "chars_per_page": float(total_chars), "warnings": []}


def _parse_txt(data: bytes) -> dict:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise _ParseWorkerError("parse_failed") from exc
    paragraphs = [" ".join(p.split()) for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        raise _ParseWorkerError("empty")
    blocks = [{"text": p, "page": 1, "size": DEFAULT_BODY_PT, "bold": False, "kind_hint": "paragraph"}
             for p in paragraphs]
    total_chars = sum(len(b["text"]) for b in blocks)
    return {"method": "text", "pages": 1, "blocks": blocks, "chars_per_page": float(total_chars), "warnings": []}


# --------------------------------------------------------------------------
# dispatch + entry point
# --------------------------------------------------------------------------

def _parse(kind: str, data: bytes, max_pages: int) -> dict:
    if not data:
        raise _ParseWorkerError("empty")
    if kind == "pdf":
        return _parse_pdf(data, max_pages)
    if kind == "docx":
        result = _parse_docx(data)
        estimated_pages = result.pop("_estimated_pages")
        if estimated_pages > max_pages:
            raise _ParseWorkerError("too_many_pages")
        return result
    if kind == "html":
        return _parse_html(data)
    if kind == "md":
        return _parse_md(data)
    if kind == "txt":
        return _parse_txt(data)
    raise _ParseWorkerError("parse_failed")


def _write(obj: dict) -> None:
    sys.stdout.buffer.write(json.dumps(obj).encode("utf-8"))
    sys.stdout.buffer.flush()


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        _write({"error": "parse_failed"})
        return 1
    kind, max_pages_arg = argv[1], argv[2]
    try:
        max_pages = int(max_pages_arg)
    except ValueError:
        _write({"error": "parse_failed"})
        return 1
    data = sys.stdin.buffer.read()
    try:
        result = _parse(kind, data, max_pages)
    except _ParseWorkerError as exc:
        _write({"error": exc.code})
        return 1
    except Exception:
        # never logs the document text/bytes, only that parsing failed; the parent discards stderr in production
        # (stderr=DEVNULL) — this is for local debugging only.
        logger.exception("unhandled error parsing kind=%s", kind)
        _write({"error": "parse_failed"})
        return 1
    _write(result)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
